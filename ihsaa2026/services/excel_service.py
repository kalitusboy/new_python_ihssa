"""خدمة Excel — استيراد وتصدير بصيغ مطابقة لـ lib/services/excel_service.dart.
تستعمل openpyxl (xlsx فقط — كافٍ ومتوافق مع Flutter excel package).
"""
from __future__ import annotations
import csv
import re
from pathlib import Path
from typing import List, Optional, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..config import DOCS_DIR
from ..models.beneficiary import Beneficiary


# ──────────────────────────────────────────────
def _parse_full_name(full: str) -> tuple[str, str]:
    full = (full or "").strip()
    if not full:
        return "", ""
    parts = re.split(r"\s+", full)
    if len(parts) == 1:
        return parts[0], ""
    if len(parts) == 2:
        return parts[0], parts[1]
    return parts[0], " ".join(parts[1:])


def _extract_birth_date(text: str) -> str:
    if not text:
        return ""
    text = text.strip()
    m = re.search(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", text)
    if m:
        return f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
    m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", text)
    if m:
        return f"{m.group(3)}-{m.group(2).zfill(2)}-{m.group(1).zfill(2)}"
    m = re.search(r"عام\s*(\d{4})", text)
    if m:
        return f"{m.group(1)}-01-01"
    m = re.match(r"^(\d{4})$", text)
    if m:
        return f"{m.group(1)}-01-01"
    return text[:50]


def _extract_birth_place(text: str) -> str:
    if not text:
        return ""
    s = text
    s = re.sub(r"\d{4}[/-]\d{1,2}[/-]\d{1,2}", "", s)
    s = re.sub(r"\d{1,2}[/-]\d{1,2}[/-]\d{4}", "", s)
    s = re.sub(r"عام\s*\d{4}", "", s)
    s = re.sub(r"\d{4}", "", s)
    s = s.strip().strip(",").strip()
    return s


# ──────────────────────────────────────────────
class ExcelService:
    # ── الاستيراد ──────────────────────────────
    def import_from_excel(self, path: Path) -> List[Beneficiary]:
        path = Path(path)
        if path.suffix.lower() == ".csv":
            return self._import_csv(path)

        wb = load_workbook(path, read_only=True, data_only=True)
        out: List[Beneficiary] = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue
            # نتجاوز صف العنوان
            for row in rows[1:]:
                cells = list(row) + [None] * 10
                full = (cells[0] or "")
                program = (cells[1] or "")
                address = (cells[2] or "")
                birth_full = (cells[3] or "")
                birth_place = (cells[4] or "")
                full = str(full).strip() if full else ""
                program = str(program).strip() if program else ""
                address = str(address).strip() if address else ""
                birth_full = str(birth_full).strip() if birth_full else ""
                birth_place = str(birth_place).strip() if birth_place else ""

                if not full:
                    continue
                fn, ln = _parse_full_name(full)
                bd = _extract_birth_date(birth_full)
                bp = birth_place or _extract_birth_place(birth_full)

                if not (fn or ln):
                    continue
                out.append(Beneficiary(
                    first_name=fn,
                    last_name=ln,
                    full_name=full,
                    birth_date=bd or None,
                    birth_place=bp or None,
                    address=address or None,
                    program=program or "عام",
                ))
        return out

    def _import_csv(self, path: Path) -> List[Beneficiary]:
        out: List[Beneficiary] = []
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            rows = list(reader)
        if len(rows) < 2:
            return out
        for row in rows[1:]:
            cells = list(row) + [""] * 10
            full = cells[0].strip()
            program = cells[1].strip()
            address = cells[2].strip()
            birth_full = cells[3].strip()
            birth_place = cells[4].strip()
            if not full:
                continue
            fn, ln = _parse_full_name(full)
            bd = _extract_birth_date(birth_full)
            bp = birth_place or _extract_birth_place(birth_full)
            if not (fn or ln):
                continue
            out.append(Beneficiary(
                first_name=fn, last_name=ln, full_name=full,
                birth_date=bd or None, birth_place=bp or None,
                address=address or None, program=program or "عام",
            ))
        return out

    # ── تصدير قائمة المستفيدين ─────────────────
    def export_to_excel(
        self,
        beneficiaries: Sequence[Beneficiary],
        file_path: Optional[Path] = None,
        sheet_title: str = "المستفيدين",
    ) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title
        try:
            ws.sheet_view.rightToLeft = True
        except Exception:
            pass

        headers = [
            "الإسم واللقب", "البرنامج", "العنوان", "تاريخ الميلاد", "مكان الميلاد",
            "كهرباء", "غاز", "مياه", "تطهير", "الحالة",
        ]
        ws.append(headers)
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="0D47A1")
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for b in beneficiaries:
            ws.append([
                b.display_name,
                b.program or "",
                b.address or "",
                b.birth_date or "",
                b.birth_place or "",
                int(b.electricity), int(b.gas), int(b.water), int(b.sewage),
                b.status,
            ])

        # ضبط أعرض تقريبي
        widths = [28, 18, 22, 14, 18, 8, 8, 8, 8, 22]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

        if file_path is None:
            from datetime import datetime
            file_path = DOCS_DIR / f"تصدير_المستفيدين_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(file_path)
        return Path(file_path)

    # ── تصدير الإحصائيات (ورقتان) ──────────────
    def export_statistics_to_file(
        self,
        file_path: Path,
        main_headers: List[str],
        main_rows: List[List],
        detail_headers: List[str],
        detail_rows: List[List],
    ) -> Path:
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "الإحصائيات العامة"
        try:
            ws1.sheet_view.rightToLeft = True
        except Exception:
            pass
        self._write_sheet(ws1, main_headers, main_rows)

        ws2 = wb.create_sheet("تفاصيل المنتهية والمشغولة")
        try:
            ws2.sheet_view.rightToLeft = True
        except Exception:
            pass
        if detail_rows:
            self._write_sheet(ws2, detail_headers, detail_rows)
        else:
            ws2.cell(1, 1, "لا توجد بيانات")

        wb.save(file_path)
        return Path(file_path)

    def _write_sheet(self, ws, headers: List[str], rows: List[List]) -> None:
        ws.append(headers)
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="0D47A1")
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in rows:
            ws.append(["" if v is None else v for v in row])
        # اضبط عرض الأعمدة تلقائياً
        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
                14, min(40, len(str(headers[col_idx - 1])) + 6)
            )
