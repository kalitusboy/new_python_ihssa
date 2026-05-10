"""خدمة توليد تقارير Word (.docx) — مكافئة لـ advanced_report_service.dart.

تُولِّد:
  1. محضر معاينة لبرنامج واحد  (export_program_report_docx)
  2. تقرير الإحصائيات الشامل   (export_advanced_stats_docx)

تدعم RTL عربي بشكل كامل عبر python-docx + إعداد الفقرات.
"""
from __future__ import annotations

import time
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.util import Twips

from ..config import DOCS_DIR
from .database_service import DatabaseService


# ─────────────────────────────────────────────────────────────────
# ألوان مطابقة لتقارير Flutter
# ─────────────────────────────────────────────────────────────────
_BLUE   = RGBColor(0x0D, 0x47, 0xA1)   # PRIMARY
_DARK   = RGBColor(0x0A, 0x38, 0x80)
_GREEN  = RGBColor(0x1B, 0x5E, 0x20)
_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
_GRAY   = RGBColor(0xF1, 0xF5, 0xF9)
_BORDER = RGBColor(0xCB, 0xD5, 0xE1)


def _set_rtl_paragraph(paragraph) -> None:
    """يجعل الفقرة RTL عربي."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "right")
    pPr.append(jc)


def _set_cell_rtl(cell) -> None:
    """يجعل خلية الجدول RTL."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBidi = OxmlElement("w:bidi")
    tcPr.append(tcBidi)
    for para in cell.paragraphs:
        _set_rtl_paragraph(para)


def _set_cell_bg(cell, hex_color: str) -> None:
    """يضبط لون خلفية خلية جدول."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)


def _set_doc_rtl(doc: Document) -> None:
    """يضبط اتجاه المستند كله RTL."""
    settings = doc.settings.element
    bidi = OxmlElement("w:bidi")
    settings.append(bidi)


def _add_heading(doc: Document, text: str, level: int = 1,
                 color: RGBColor = _BLUE) -> None:
    p = doc.add_paragraph()
    _set_rtl_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16 if level == 1 else 13)
    run.font.color.rgb = color
    run.font.name = "Cairo"
    p.space_after = Pt(6)


def _add_subheading(doc: Document, text: str) -> None:
    _add_heading(doc, text, level=2, color=_DARK)


def _add_line(doc: Document, text: str, bold: bool = False, size: int = 11) -> None:
    p = doc.add_paragraph()
    _set_rtl_paragraph(p)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Cairo"


def _build_table(doc: Document, headers: List[str], rows: List[List],
                 header_bg: str = "0D47A1") -> None:
    """يبني جدول مع رأس ملوّن وتوسيط RTL."""
    col_count = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # ── الرأس ──────────────────────────────────
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr_cells[i]
        _set_cell_bg(cell, header_bg)
        _set_cell_rtl(cell)
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(h)
            run.bold = True
            run.font.color.rgb = _WHITE
            run.font.size = Pt(11)
            run.font.name = "Cairo"

    # ── البيانات ───────────────────────────────
    for ri, row_data in enumerate(rows):
        row_cells = table.rows[ri + 1].cells
        bg = "F1F5F9" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row_data):
            cell = row_cells[ci]
            _set_cell_bg(cell, bg)
            _set_cell_rtl(cell)
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(str(val))
                run.font.size = Pt(10)
                run.font.name = "Cairo"

    doc.add_paragraph()  # مسافة بعد الجدول


# ─────────────────────────────────────────────────────────────────
class WordReportService:
    """توليد تقارير Word — Singleton."""

    _instance: "WordReportService" = None  # type: ignore

    def __new__(cls):
        if cls._instance is None:
            cls._instance = super().__new__(cls)
        return cls._instance

    def __init__(self):
        self.db = DatabaseService()

    # ══════════════════════════════════════════════════════════════
    # ① محضر معاينة برنامج واحد
    # ══════════════════════════════════════════════════════════════
    def export_program_report_docx(
        self,
        program: str,
        wilaya: str = "",
        daira: str = "",
        baladia: str = "",
        place: str = "",
        num: str = "",
        members: Optional[List[dict]] = None,
        output_path: Optional[Path] = None,
    ) -> Path:
        """يولِّد محضر معاينة لبرنامج (مطابق لـ report_screen في Flutter)."""
        if output_path is None:
            ts = int(time.time())
            safe = program.replace("/", "_").replace("\\", "_")[:40]
            output_path = DOCS_DIR / f"محضر_{safe}_{ts}.docx"

        stats  = self.db.get_report_stats(program)
        images = self.db.get_program_images(program)
        now    = datetime.now().strftime("%Y-%m-%d")

        doc = Document()
        _set_doc_rtl(doc)

        # إعدادات الصفحة
        section = doc.sections[0]
        section.right_margin = Cm(2)
        section.left_margin  = Cm(2)
        section.top_margin   = Cm(1.5)
        section.bottom_margin = Cm(1.5)

        # ── الترويسة ──────────────────────────────────────────────
        _add_heading(doc, "الجمهورية الجزائرية الديمقراطية الشعبية")
        _add_heading(doc, "ولاية النعامة — مديرية السكن والتجهيزات العمومية")
        _add_heading(doc, f"محضر معاينة — برنامج: {program}", color=_DARK)
        _add_line(doc, f"التاريخ: {now}  |  المكان: {place or '___'}", size=10)
        doc.add_paragraph()

        # ── معلومات الوثيقة ───────────────────────────────────────
        if wilaya or daira or baladia:
            _build_table(doc,
                headers=["الولاية", "الدائرة", "البلدية", "رقم الوثيقة"],
                rows=[[wilaya, daira, baladia, num]],
                header_bg="0D47A1"
            )

        # ── الحالة العامة ─────────────────────────────────────────
        _add_subheading(doc, "الحالة العامة")
        _build_table(doc,
            headers=["البيان", "العدد"],
            rows=[
                ["الحصة (الإجمالي)",       str(stats.get("quota", 0))],
                ["المُحصاة",               str(stats.get("done", 0))],
                ["في طور الانجاز",         str(stats.get("in_progress", 0))],
                ["على مستوى الأعمدة",      str(stats.get("pillars", 0))],
                ["منتهية غير مشغولة",      str(stats.get("finished_not_occupied", 0))],
                ["منتهية ومشغولة",         str(stats.get("finished_occupied", 0))],
            ]
        )

        # ── الشبكات ───────────────────────────────────────────────
        _add_subheading(doc, "الربط بالشبكات (للمنتهية والمشغولة)")
        _build_table(doc,
            headers=["الشبكة", "العدد"],
            rows=[
                ["كهرباء",             str(stats.get("elec_occ", 0))],
                ["غاز",               str(stats.get("gas_occ", 0))],
                ["مياه",              str(stats.get("water_occ", 0))],
                ["تطهير",             str(stats.get("sew_occ", 0))],
                ["جميع الشبكات معاً", str(stats.get("fully_connected", 0))],
            ]
        )

        # ── لجنة المعاينة ─────────────────────────────────────────
        if members:
            _add_subheading(doc, "أعضاء لجنة المعاينة")
            member_rows = [[m.get("name", ""), m.get("role", "")] for m in members]
            _build_table(doc, headers=["الاسم واللقب", "الصفة"],
                         rows=member_rows)

        # ── قائمة المستفيدين (بصور) ──────────────────────────────
        if images:
            _add_subheading(doc, f"قائمة المستفيدين بصور ({len(images)} مستفيد)")
            img_rows = [
                [str(i + 1), f"{img.get('first_name','')} {img.get('last_name','')}".strip()]
                for i, img in enumerate(images)
            ]
            _build_table(doc, headers=["#", "الاسم"], rows=img_rows)

        # ── التوقيعات ─────────────────────────────────────────────
        doc.add_paragraph()
        _add_subheading(doc, "التوقيعات")
        sig_table = doc.add_table(rows=2, cols=3)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        labels = ["رئيس اللجنة", "ممثل البلدية", "ممثل المستفيدين"]
        for i, lbl in enumerate(labels):
            cell = sig_table.rows[0].cells[i]
            _set_cell_rtl(cell)
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(lbl)
                run.bold = True
                run.font.name = "Cairo"
            # خلية التوقيع
            sig_cell = sig_table.rows[1].cells[i]
            _set_cell_rtl(sig_cell)
            for para in sig_cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.add_run("\n\n\n")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    # ══════════════════════════════════════════════════════════════
    # ② تقرير الإحصائيات الشامل
    # ══════════════════════════════════════════════════════════════
    def export_advanced_stats_docx(
        self, output_path: Optional[Path] = None
    ) -> Path:
        """تقرير الإحصائيات المتقدمة (مطابق لـ advanced_report_service.dart)."""
        if output_path is None:
            ts = int(time.time())
            output_path = DOCS_DIR / f"تقرير_شامل_{ts}.docx"

        adv = self.db.get_advanced_stats()
        tot = adv.get("totals") or {}
        now = datetime.now().strftime("%Y-%m-%d %H:%M")

        doc = Document()
        _set_doc_rtl(doc)

        section = doc.sections[0]
        section.right_margin = Cm(1.5)
        section.left_margin  = Cm(1.5)
        section.top_margin   = Cm(1.5)
        section.bottom_margin = Cm(1.5)

        # ── الترويسة ──────────────────────────────────────────────
        _add_heading(doc, "التقرير الإحصائي المتقدم")
        _add_heading(doc, "إحصاء السكن الريفي 2026 — نسيم الحوضان",
                     level=2, color=_DARK)
        _add_line(doc, f"تاريخ التوليد: {now}", size=10)
        doc.add_paragraph()

        # ── الملخص العام ──────────────────────────────────────────
        _add_subheading(doc, "الملخص العام")
        _build_table(doc,
            headers=["البيان", "العدد"],
            rows=[
                ["الإجمالي",           str(tot.get("total", 0))],
                ["المحصاة",           str(tot.get("done", 0))],
                ["بصور",              str(tot.get("with_image", 0))],
                ["بدون صور",          str(tot.get("without_image", 0))],
                ["كهرباء",            str(tot.get("elec", 0))],
                ["غاز",              str(tot.get("gas", 0))],
                ["مياه",             str(tot.get("water", 0))],
                ["تطهير",            str(tot.get("sewage", 0))],
            ]
        )

        # ── حسب البرامج ───────────────────────────────────────────
        _add_subheading(doc, "الإحصائيات حسب البرنامج")
        prog_headers = [
            "البرنامج", "الحصة", "المحصاة",
            "في طور الانجاز", "على الأعمدة",
            "منتهية غ.مشغولة", "منتهية مشغولة",
            "كهرباء", "غاز", "مياه", "تطهير", "بصور",
        ]
        prog_rows = []
        for r in adv.get("byProgram") or []:
            prog_rows.append([
                r.get("program") or "",
                r.get("total", 0),    r.get("done", 0),
                r.get("s1", 0),       r.get("s2", 0),
                r.get("s3", 0),       r.get("s4", 0),
                r.get("elec", 0),     r.get("gas", 0),
                r.get("water", 0),    r.get("sewage", 0),
                r.get("with_image", 0),
            ])
        _build_table(doc, prog_headers, prog_rows, header_bg="0D47A1")

        # ── حسب الحالة ────────────────────────────────────────────
        _add_subheading(doc, "الإحصائيات حسب الحالة")
        status_headers = [
            "الحالة", "العدد", "كهرباء", "غاز",
            "مياه", "تطهير", "بدون شبكة", "بصور",
        ]
        status_rows = []
        for r in adv.get("byStatus") or []:
            status_rows.append([
                r.get("status") or "",
                r.get("total", 0),  r.get("elec", 0),  r.get("gas", 0),
                r.get("water", 0),  r.get("sewage", 0), r.get("none", 0),
                r.get("with_image", 0),
            ])
        _build_table(doc, status_headers, status_rows, header_bg="0A3880")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    # ══════════════════════════════════════════════════════════════
    # ③ تصدير صور برنامج كـ ZIP
    # ══════════════════════════════════════════════════════════════
    def export_program_images_zip(
        self, program: str, output_path: Optional[Path] = None
    ) -> Path:
        """يضغط صور برنامج في ZIP (مطابق لـ getSafeZip في Flutter)."""
        if output_path is None:
            ts = int(time.time())
            safe = program.replace("/", "_").replace("\\", "_")[:30]
            output_path = DOCS_DIR / f"صور_{safe}_{ts}.zip"

        images = self.db.get_program_images(program)
        if not images:
            raise Exception(f"لا توجد صور للبرنامج: {program}")

        import zipfile
        added = 0
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as z:
            seen: set[str] = set()
            for img in images:
                p = Path(img.get("path") or "")
                if not p.exists():
                    continue
                name = img.get("name") or p.name
                # تفادي التكرار
                base, ext = Path(name).stem, Path(name).suffix
                unique = name
                counter = 1
                while unique in seen:
                    unique = f"{base}_{counter}{ext}"
                    counter += 1
                seen.add(unique)
                z.write(p, arcname=unique)
                added += 1

        if added == 0:
            output_path.unlink(missing_ok=True)
            raise Exception("لا توجد ملفات صور موجودة على القرص")

        return output_path
